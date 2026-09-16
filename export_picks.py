"""
RoachFinder 3.0 — Excel export of weekly picks

Writes data/current_week.json into picks/RoachFinder_Picks_<season>.xlsx,
one sheet per week ("Week 02", "Week 03", ...). Re-running a week replaces
that week's sheet; other weeks are kept.

Called automatically at the end of run_week.py. Can also be run by hand:
    python export_picks.py
"""

import json
import os

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from config import PROJECT_ROOT, SEASON_LABEL, data_path

PICKS_DIR = os.path.join(PROJECT_ROOT, "picks")
PICKS_FILE = os.path.join(PICKS_DIR, f"RoachFinder_Picks_{SEASON_LABEL}.xlsx")

COLUMNS = [
    ("Rank", 6), ("Game Time", 11), ("Away", 16), ("Home", 16),
    ("Vegas Line", 12), ("O/U", 7), ("Pick", 18), ("Confidence", 11),
    ("Top 5", 7), ("Above Median", 13), ("Result (W/L/P)", 15),
]


def export_week_to_excel():
    with open(data_path("current_week.json"), "r") as f:
        week_data = json.load(f)

    week = week_data["week"]
    top5 = set(week_data.get("top_5_picks", []))
    above = set(week_data.get("above_median_picks", []))
    games = sorted(week_data["games"], key=lambda g: g["confidence_factor"], reverse=True)

    os.makedirs(PICKS_DIR, exist_ok=True)
    wb = load_workbook(PICKS_FILE) if os.path.exists(PICKS_FILE) else Workbook()
    if "Sheet" in wb.sheetnames and len(wb.sheetnames) == 1:
        del wb["Sheet"]  # default empty sheet from a new workbook

    sheet_name = f"Week {week:02d}"
    if sheet_name in wb.sheetnames:
        del wb[sheet_name]

    # Keep sheets in week order
    later = [n for n in wb.sheetnames if n.startswith("Week ") and n > sheet_name]
    index = wb.sheetnames.index(later[0]) if later else len(wb.sheetnames)
    ws = wb.create_sheet(sheet_name, index)

    header_fill = PatternFill("solid", fgColor="1F2937")
    for col, (name, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col, value=name)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
        ws.column_dimensions[cell.column_letter].width = width

    for rank, g in enumerate(games, start=1):
        label = g["winning_pick_label"]
        ou = g.get("over_under")
        try:
            ou = float(ou)
        except (TypeError, ValueError):
            pass
        ws.append([
            rank,
            g.get("game_time", ""),
            g["away_team"],
            g["home_team"],
            g.get("spread_line", ""),
            ou,
            label,
            round(g["confidence_factor"], 2),
            "Yes" if label in top5 else "",
            "Yes" if label in above else "",
            "",
        ])

    footer = len(games) + 3
    ws.cell(row=footer, column=1, value="Rolling median CF:").font = Font(bold=True)
    ws.cell(row=footer, column=4, value=week_data.get("rolling_median_cf"))
    ws.freeze_panes = "A2"

    try:
        wb.save(PICKS_FILE)
    except PermissionError:
        print(f"\n  WARNING: Could not write {PICKS_FILE}.")
        print("  Close it in Excel, then run: python export_picks.py")
        return
    print(f"\n  Saved Week {week} picks to {PICKS_FILE} (sheet '{sheet_name}')")


if __name__ == "__main__":
    export_week_to_excel()
